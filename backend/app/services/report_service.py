"""CSV/Excel export — mirrors the deck 3-sheet layout in local-only mode."""

from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def _sanitize(value: str) -> str:
    if value and value[0] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def build_csv(record) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(["Steelshield — synthetic report", record.claim_level])
    writer.writerow(["run_id", record.id, "target", record.target])
    writer.writerow(
        [
            "mode",
            getattr(record, "mode", "fixture"),
            "attacker",
            getattr(record, "attacker", "replay"),
        ]
    )
    writer.writerow(
        [
            "status",
            getattr(record, "status", "completed"),
            "egress_host",
            getattr(record, "egress_host", "") or "",
        ]
    )
    writer.writerow(["split", record.split, "policies", ",".join(record.policy_ids)])
    writer.writerow(
        ["review_status", record.review_status, "manifest_version", record.manifest_version]
    )
    writer.writerow(
        [
            "external_requests",
            record.metrics.external_requests,
            "external_bytes",
            record.metrics.external_bytes,
        ]
    )
    writer.writerow([])
    writer.writerow(
        [
            "case_id",
            "family_id",
            "type",
            "policy",
            "suite",
            "severity",
            "decision",
            "expected",
            "expectation_met",
        ]
    )
    for row in record.case_results:
        writer.writerow(
            [
                _sanitize(row["case_id"]),
                _sanitize(row["family_id"]),
                row["case_type"],
                row["policy_id"],
                row["suite"],
                row["severity"],
                row["decision"],
                row["expected"],
                row["expectation_met"],
            ]
        )
    writer.writerow([])
    writer.writerow(["finding_case_id", "title", "severity", "decision", "explanation"])
    for finding in record.findings:
        writer.writerow(
            [
                _sanitize(finding.case_id),
                _sanitize(finding.title),
                finding.severity,
                finding.decision,
                _sanitize(finding.explanation),
            ]
        )
    return buf.getvalue()


def build_xlsx(record) -> bytes:
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "summary"
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    ws_summary.append([f"Steelshield — synthetic report — {record.claim_level}"])
    ws_summary.append(["run_id", record.id, "target", record.target])
    ws_summary.append(
        [
            "mode",
            getattr(record, "mode", "fixture"),
            "attacker",
            getattr(record, "attacker", "replay"),
        ]
    )
    ws_summary.append(
        [
            "status",
            getattr(record, "status", "completed"),
            "egress_host",
            getattr(record, "egress_host", "") or "",
        ]
    )
    ws_summary.append(["split", record.split, "policies", ", ".join(record.policy_ids)])
    ws_summary.append(
        ["review_status", record.review_status, "manifest_version", record.manifest_version]
    )
    ws_summary.append(
        [
            "external_requests",
            record.metrics.external_requests,
            "external_bytes",
            record.metrics.external_bytes,
        ]
    )
    ws_summary.append(
        [
            "total_cases",
            record.metrics.total_cases,
            "attack",
            record.metrics.attack_cases,
            "benign",
            record.metrics.benign_cases,
        ]
    )
    ws_summary.append(
        [
            "precision",
            record.metrics.precision,
            "recall",
            record.metrics.recall,
            "FPR",
            record.metrics.false_positive_rate,
        ]
    )

    ws_cases = wb.create_sheet("cases")
    ws_cases.append(
        [
            "case_id",
            "family_id",
            "type",
            "policy",
            "suite",
            "severity",
            "decision",
            "expected",
            "expectation_met",
        ]
    )
    for cell in ws_cases[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in record.case_results:
        ws_cases.append(
            [
                _sanitize(row["case_id"]),
                _sanitize(row["family_id"]),
                row["case_type"],
                row["policy_id"],
                row["suite"],
                row["severity"],
                row["decision"],
                row["expected"],
                row["expectation_met"],
            ]
        )

    ws_findings = wb.create_sheet("findings")
    ws_findings.append(
        [
            "case_id",
            "title",
            "severity",
            "decision",
            "explanation",
            "oracle_leaked",
            "prompt_flagged",
        ]
    )
    for cell in ws_findings[1]:
        cell.fill = header_fill
        cell.font = header_font
    for finding in record.findings:
        ws_findings.append(
            [
                _sanitize(finding.case_id),
                _sanitize(finding.title),
                finding.severity,
                finding.decision,
                _sanitize(finding.explanation),
                finding.evidence["oracle"]["leaked"],
                finding.evidence["prompt_screen"]["flagged"],
            ]
        )
    for ws in (ws_summary, ws_cases, ws_findings):
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
